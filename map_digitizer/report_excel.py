"""
Stage 10 — Excel deliverable.

Builds the same workbook shape as the worked examples (Read Me sheet with
methodology/caveats up front, then one sheet per finding category) but
parameterized by feature_type so the same builder produces sensible sheet
names and columns whether you're digitizing fields, facilities, or pipelines.
"""

from __future__ import annotations

from typing import Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from .config import MapConfig, FeatureType

_NOUN = {
    FeatureType.AREA: ("field", "fields", "polygon", "polygons", "outline"),
    FeatureType.POINT: ("facility", "facilities", "point", "points", "location"),
    FeatureType.LINE: ("pipeline segment", "pipeline segments", "line", "lines", "route"),
}


def _nouns(feature_type: FeatureType):
    return _NOUN.get(feature_type, ("feature", "features", "shape", "shapes", "outline"))


def _write_readme(wb: Workbook, title: str, paragraphs: list) -> None:
    ws = wb.active
    ws.title = "Read Me"
    ws.column_dimensions["A"].width = 120
    row = 1
    ws.cell(row=row, column=1, value=title).font = Font(bold=True, size=14)
    row += 2
    for heading, body in paragraphs:
        if heading:
            ws.cell(row=row, column=1, value=heading).font = Font(bold=True, size=11)
            row += 1
        for line in body.split("\n"):
            cell = ws.cell(row=row, column=1, value=line)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            row += 1
        row += 1


def _write_dataframe_sheet(wb: Workbook, sheet_name: str, df: pd.DataFrame, unverified_cols: Optional[list] = None) -> None:
    ws = wb.create_sheet(sheet_name)
    unverified_cols = unverified_cols or []
    headers = list(df.columns)
    for col_idx, header in enumerate(headers, start=1):
        label = f"{header} (unverified / OCR)" if header in unverified_cols else header
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = Font(bold=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(60, max(12, len(str(label)) + 2))

    for row_idx, row in enumerate(df.itertuples(index=False), start=2):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)


def build_extraction_workbook(config: MapConfig, extracted_df: pd.DataFrame,
                               georef_summary: dict, method_notes: str,
                               out_path: str, name_column: str = "candidate_name") -> str:
    """
    Standalone extraction deliverable (no crosscheck yet) — mirrors
    example_1_raw_extraction.xlsx.
    """
    singular, plural, shape_sg, shape_pl, outline_word = _nouns(config.feature_type)
    wb = Workbook()

    readme_paragraphs = [
        ("Purpose",
         f"This workbook contains {plural} extracted from {config.name}, georeferenced to real-world "
         f"WGS84 coordinates. Produced by the map_digitizer pipeline (config-driven, no AI extraction "
         f"calls) rather than manual digitizing."),
        ("Method", method_notes),
        ("Georeferencing",
         f"Fit degree: {georef_summary.get('degree')}. Control points used: {georef_summary.get('n_used')} "
         f"(dropped as outliers: {georef_summary.get('n_dropped')}). Mean residual: "
         f"{georef_summary.get('mean_residual_km'):.1f} km. Worst-case residual: "
         f"{georef_summary.get('max_residual_km'):.1f} km across control points. This is a rough accuracy "
         f"bound, not uniform across the whole map — corners and areas with fewer/weaker control points "
         f"can be meaningfully worse. Good for identifying and roughly locating {plural}, not survey-grade "
         f"boundaries."),
        ("What's in this workbook",
         f"'Extracted {plural}' lists every {singular} kept after filtering, with its feature class, "
         f"candidate name (OCR, unverified unless noted), approximate centroid coordinates, relative size, "
         f"label match distance, and the traced {outline_word} as WGS84 WKT."),
        ("Limitations",
         f"Pixel-traced {shape_pl} are approximate, not survey-grade. OCR-derived names will contain "
         f"errors and require a human pass before being treated as authoritative. This is a triage tool."),
    ]
    _write_readme(wb, f"{config.name} — {plural.title()} Extraction", readme_paragraphs)
    _write_dataframe_sheet(wb, f"Extracted {plural}".title()[:31], extracted_df, unverified_cols=[name_column])

    wb.save(out_path)
    return out_path


def build_crosscheck_workbook(config: MapConfig, extracted_df: pd.DataFrame, full_match_df: pd.DataFrame,
                               missing_from_tracker_df: pd.DataFrame, missing_geometry_df: pd.DataFrame,
                               headline_notes: str, out_path: str, name_column: str = "candidate_name") -> str:
    """
    Crosscheck deliverable against the tracker database — mirrors
    example_2_goget_crosscheck.xlsx's three-sheet structure.
    """
    singular, plural, shape_sg, shape_pl, outline_word = _nouns(config.feature_type)
    wb = Workbook()

    readme_paragraphs = [
        ("Purpose",
         f"This workbook crosschecks the tracker database against {len(extracted_df)} {plural} extracted "
         f"from {config.name}. It flags (a) map {plural} with no confident match in the tracker export, "
         f"and (b) tracker records that have no {outline_word} (WKT) data, with a candidate {outline_word} "
         f"suggested wherever a plausible match was found."),
        ("Headline numbers", headline_notes),
        ("How matching works",
         f"Each map {singular}'s centroid was compared to every tracker record's lat/lon (haversine "
         f"distance), and separately, the {singular}'s OCR candidate name was fuzzy-matched (rapidfuzz "
         f"token_sort_ratio) against tracker records within {config.crosscheck.name_match_search_cap_km:.0f} km. "
         f"Categories: 'contained/very close' (<= {config.crosscheck.contained_distance_km:.0f} km), "
         f"'name match' (fuzzy score >= {config.crosscheck.name_score_threshold:.0f}), "
         f"'nearby' (<= {config.crosscheck.nearby_distance_km:.0f} km, weak/no name signal), "
         f"'no confident match' (neither)."),
        ("Important caveats",
         "A tracker database that only carries named, notable records will always have fewer entries than "
         "a map that plots every small sub-feature — 'no confident match' is a mix of genuine gaps and "
         "features the tracker was never expected to carry. OCR misreads are usually the single biggest "
         "cause of a real match failing to register by name — always check the raw closest-distance number "
         "even for 'no match' rows. Treat this sheet as a prioritized starting point for manual review, not "
         "a final answer."),
    ]
    _write_readme(wb, f"{config.name} — Tracker Crosscheck", readme_paragraphs)

    _write_dataframe_sheet(wb, f"Candidate missing {plural}"[:31], missing_from_tracker_df, unverified_cols=[name_column])
    _write_dataframe_sheet(wb, f"Candidate missing {outline_word}s"[:31], missing_geometry_df)
    _write_dataframe_sheet(wb, "Full match table (audit)"[:31], full_match_df, unverified_cols=[name_column])

    wb.save(out_path)
    return out_path
